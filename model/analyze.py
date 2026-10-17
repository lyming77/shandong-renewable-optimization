# -*- coding: utf-8 -*-
"""
山东电网新能源发电能力优化试运行方案 —— 场站影响测算模型

依据《山东电网新能源发电能力优化试运行方案》(PDF) 与附件2参考算例实现：
  1. 触发条件：出清计划值 > 额定容量×10%，且连续2个出清周期实际出力低于计划、
     偏差比例 (P计划-P实际)/P计划 超过触发设定比例；
  2. 优化逻辑：ΔP% = β × [(P计划,-1-P实际,-1)/P计划,-1 + (P计划,-2-P实际,-2)/P计划,-2]/2，
     对触发后(跳过1个计算周期)连续α个周期的超短期预测按下式压减：
     P优化,i = P原上报,i × (1 - ΔP%)；
  3. 恢复逻辑：连续2周期"执行到位"(实际≥计划)后，分γ个周期恢复
     (第1次恢复ΔP%/2，第2次完全恢复；压减窗口内的数据不恢复)。

模型口径(可在此调整)：
  - 出清周期 = 15 分钟(与数据一致)
  - 计划基准：未优化时用数据"实时出清计划功率"列；优化窗口内用"超短期预测×(1-ΔP%)"
  - 压减对象：超短期预测功率(第2点)，压减量即"出力/电量影响"
"""
import os
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ---------------- 参数配置 ----------------
DATA_PATH = os.path.join(os.path.dirname(__file__), "..", "测算数据.xlsx")
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")

DT_HOUR = 15 / 60.0          # 单个出清周期时长(小时)
DELAY = 1                    # 触发后跳过的计算周期数(附件算例: T1/T2触发->T4起优化)
RATED_RATIO = 0.10           # 触发门槛: 计划值 > 额定容量×该比例
TRACK_TOL = 1e-6             # "执行到位"判定容差(实际 >= 计划 - 容差)

# 七个试运行场景
SCENARIOS = {
    1: dict(trigger_ratio=0.10, alpha=3, beta=1.0, gamma=2),
    2: dict(trigger_ratio=0.10, alpha=3, beta=1.2, gamma=2),
    3: dict(trigger_ratio=0.10, alpha=4, beta=1.0, gamma=2),
    4: dict(trigger_ratio=0.10, alpha=4, beta=1.3, gamma=1),
    5: dict(trigger_ratio=0.10, alpha=4, beta=1.5, gamma=2),
    6: dict(trigger_ratio=0.15, alpha=4, beta=1.5, gamma=2),
    7: dict(trigger_ratio=0.15, alpha=5, beta=1.5, gamma=2),
}

# 数据Sheet -> (实际场站名, 装机容量MW)。玉山=临沭龙源, 东盘=玉山龙源(据用户提供的容量表)
STATIONS = {
    "北海": ("北海龙源", 75.60),
    "玉山": ("临沭龙源", 98.20),
    "陈公": ("陈公龙源", 110.00),
    "滨海": ("滨海龙源", 99.00),
    "东盘": ("玉山龙源", 76.50),
    "冯家": ("冯家龙源", 95.00),
    "阜山": ("阜山龙源", 479.64),
    "六合": ("六合龙源", 299.84),
}


# ---------------- 数据读取 ----------------
def _to_num(v):
    """把单元格值转 float：数字直接转，数字文本(带引号)也转，空/异常转 NaN。"""
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().strip("'").strip('"').strip()
        if s == "":
            return np.nan
        try:
            return float(s)
        except ValueError:
            return np.nan
    return np.nan


def _ffill(arr):
    mask = np.isnan(arr)
    idx = np.where(~mask, np.arange(len(arr)), 0)
    np.maximum.accumulate(idx, out=idx)
    return arr[idx]


def _bfill(arr):
    n = len(arr)
    mask = np.isnan(arr)
    idx = np.where(~mask, np.arange(n), n - 1)
    np.minimum.accumulate(idx[::-1], out=idx[::-1])
    return arr[idx]


def _read_sheet(ws):
    """把单个工作表(openpyxl ws)解析为各列数组及时间标签。"""
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]  # 去掉表头

    def col(i):
        return np.array([_to_num(r[i]) for r in data], dtype=float)

    ts = []
    for r in data:
        y, m, d, t = r[0], r[1], r[2], str(r[3]).strip().strip("'")
        ts.append(f"{int(y)}-{int(m):02d}-{int(d):02d} {t}")
    ts = np.array(ts)

    d = {
        "ts": ts,
        "price": col(4),        # 实时电价
        "plan": col(5),         # 实时出清出力计划功率
        "forecast": col(6),     # 超短期第二个点预测功率(优化对象)
        "da_forecast": col(7),  # 日前预测
        "actual": col(8),       # 实际功率
        "avail": col(9),        # 理论可用功率
    }
    # 缺失值填充(前向+后向)；预测列若仍缺则回退到计划列
    for k in ("forecast", "plan", "actual", "avail", "price", "da_forecast"):
        d[k] = _bfill(_ffill(d[k]))
    d["forecast"] = np.where(np.isnan(d["forecast"]), d["plan"], d["forecast"])
    return d


def list_sheets(source=DATA_PATH):
    """列出工作簿所有 Sheet 名。"""
    wb = openpyxl.load_workbook(source, read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def load_sheets(source, sheets):
    """一次打开工作簿，读回指定 sheets -> {sheet名: data_dict}。"""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    out = {}
    for name in sheets:
        if name in wb.sheetnames:
            out[name] = _read_sheet(wb[name])
    wb.close()
    return out


def load_station(sheet, source=DATA_PATH):
    """读取单个场站Sheet(薄封装，便于单独读取)。"""
    return load_sheets(source, [sheet])[sheet]


def build_template_workbook(year=2026, month=8, stations=None):
    """生成数据采集样表：填写说明 + 每场站一个 Sheet(表头 + 整月 15 分钟时间轴，数据列留空)。"""
    import calendar
    if stations is None:
        stations = STATIONS
    wb = openpyxl.Workbook()

    # 填写说明
    ws0 = wb.active
    ws0.title = "填写说明"
    notes = [
        ["山东电网新能源发电能力优化 —— 数据采集样表"],
        [],
        ["每个场站填写一个 Sheet，Sheet 名 = 场站名（可自行增删场站 Sheet）。"],
        ["数据为 15 分钟一个点；时间列(年/月/日/时点)已按整月预填，只需填后面 6 列。"],
        [],
        ["各列含义："],
        ["实时电价", "该时点实时出清电价，单位 元/MWh（可为负）"],
        ["实时出清出力计划功率（MW）", "该时点实时市场出清计划功率"],
        ["超短期第二个点预测功率（MW）", "场站上报的超短期预测功率（优化/压减对象）"],
        ["日前预测功率", "日前预测功率"],
        ["实际功率（MW）（5分钟）", "场站实际发电功率"],
        ["理论可用功率（MW）（5分钟）", "理论可用功率（可用容量）"],
        [],
        ["提示：光伏场站夜间无出力的时点可留空，模型会按前后值自动填充。"],
    ]
    for r in notes:
        ws0.append(r)
    ws0.column_dimensions["A"].width = 34
    ws0.column_dimensions["B"].width = 60

    header = ["年", "月", "日", "时点", "实时电价", "实时出清出力计划功率（MW）",
              "超短期第二个点预测功率（MW）", "日前预测功率", "实际功率（MW）（5分钟）",
              "理论可用功率（MW）（5分钟）"]

    # 一天的 15 分钟时点(与测算数据一致: 00:15~24:00，共 96 点)
    def day_times():
        t = []
        for h in range(24):
            for m in (0, 15, 30, 45):
                if h == 0 and m == 0:
                    continue
                t.append(f"{h:02d}:{m:02d}")
        t.append("24:00")
        return t

    times = day_times()
    num_days = calendar.monthrange(year, month)[1]

    for sheet, (real_name, rated) in stations.items():
        ws = wb.create_sheet(sheet)
        ws.append(header)
        for day in range(1, num_days + 1):
            for t in times:
                ws.append([year, month, day, t, None, None, None, None, None, None])

    return wb


# ---------------- 模拟引擎 ----------------
def simulate(d, rated, sc, delay=DELAY):
    """按单一场景回放优化逻辑，返回逐周期结果与事件记录。"""
    n = len(d["forecast"])
    forecast = d["forecast"]
    plan = d["plan"]
    actual = d["actual"]
    price = d["price"]

    sched = np.zeros(n)          # 各周期预测压减比例(优化窗口内=ΔP%)
    triggers = []                # (触发时刻idx, ΔP%, dev0, dev1)
    recovers = []                # (恢复时刻idx, 恢复前ΔP%)

    i = 1
    recover_block_until = -1     # 恢复分批期间不重复判定恢复
    window_end = -1              # 当前压减窗口结束索引(压减窗口内数据不恢复)
    while i < n:
        # 当前两周期"计划值": 优化窗口内取压减后预测, 否则取历史出清计划
        p0 = forecast[i-1] * (1 - sched[i-1]) if sched[i-1] > 0 else plan[i-1]
        p1 = forecast[i] * (1 - sched[i]) if sched[i] > 0 else plan[i]
        a0, a1 = actual[i-1], actual[i]

        gate = p0 > RATED_RATIO * rated and p1 > RATED_RATIO * rated
        dev0 = (p0 - a0) / p0 if p0 > 1e-6 else 0.0
        dev1 = (p1 - a1) / p1 if p1 > 1e-6 else 0.0
        below = a0 < p0 and a1 < p1

        # --- 触发: 连续2周期偏差超比例 ---
        if gate and below and dev0 > sc["trigger_ratio"] and dev1 > sc["trigger_ratio"]:
            dP = sc["beta"] * (dev0 + dev1) / 2.0
            start = i + 1 + delay
            end = min(start + sc["alpha"], n)
            window_end = end
            if start < n:
                for k in range(start, end):
                    sched[k] = dP
            triggers.append((i, dP, dev0, dev1))
            recover_block_until = -1            # 新一次优化, 恢复判定复位
            i = max(start + 1, i + 1) if start < n else i + 1  # 跳到优化窗口内, 避免延迟周期重复触发
            continue

        # --- 恢复: 当前周期处于压减态且连续2周期执行到位(实际>=计划) ---
        if sched[i] > 0 and i > recover_block_until:
            met = a0 >= p0 - TRACK_TOL and a1 >= p1 - TRACK_TOL
            if met:
                dP = sched[i]
                if dP > 0:
                    # 压减窗口内的预测数据不恢复：恢复批次从窗口之后开始(不截断窗口)
                    rec_start = max(window_end, i + 1)
                    if sc["gamma"] >= 2:
                        if rec_start < n:
                            sched[rec_start] = dP / 2.0
                        if rec_start + 1 < n:
                            sched[rec_start + 1] = 0.0
                    elif sc["gamma"] == 1:
                        if rec_start < n:
                            sched[rec_start] = 0.0
                    recovers.append((i, dP))
                    recover_block_until = rec_start + sc["gamma"]
                    i += 1
                    continue
        i += 1

    reduced = sched > 0
    # 调整后计划值: 被调整周期=预测×(1-ΔP%), 否则=原出清计划
    opt_plan = np.where(reduced, forecast * (1 - sched), plan)
    # 压减后实际出力 = min(之前实际, 压减后计划), 未调整周期保持原实际
    new_actual = np.where(reduced, np.minimum(actual, opt_plan), actual)
    # 压减量 = 之前实际 - 压减后实际出力 = max(0, 之前实际 - 压减后计划)
    cut_power = np.where(reduced, np.maximum(0.0, actual - opt_plan), 0.0)
    cut_energy = float(np.sum(cut_power) * DT_HOUR)     # 压减电量(MWh)
    actual_energy = float(np.sum(actual) * DT_HOUR)     # 总实际电量(MWh)

    # ---- 收益分析: 压减电量(MWh) × 实时电价(元/MWh) ----
    cut_energy_t = cut_power * DT_HOUR                  # 逐周期压减电量(MWh)
    rev_t = cut_energy_t * price                        # 逐周期收益影响(元): 正价=少赚, 负价=少亏
    pos = price > 0
    neg = price < 0
    cut_energy_pos = float(np.sum(cut_energy_t[pos]))   # 正价时段压减电量
    cut_energy_neg = float(np.sum(cut_energy_t[neg]))   # 负价时段压减电量
    rev_loss_pos = float(np.sum(rev_t[pos]))            # 正价少赚(损失, 元)
    rev_save_neg = float(-np.sum(rev_t[neg]) + 0.0)     # 负价少亏(避免支出, 元, 正值; +0.0 归一化 -0.0)
    net_rev = rev_loss_pos - rev_save_neg               # 净收益影响(元): 正=净少赚, 负=净少亏

    return {
        "sched": sched,
        "opt_plan": opt_plan,
        "new_actual": new_actual,
        "triggers": triggers,
        "recovers": recovers,
        "reduced_periods": int(reduced.sum()),
        "cut_power_peak": float(np.max(cut_power)) if reduced.any() else 0.0,
        "cut_energy": cut_energy,
        "cut_energy_pct": cut_energy / actual_energy * 100 if actual_energy > 0 else 0.0,
        "avg_dp": float(np.mean(sched[reduced])) if reduced.any() else 0.0,
        "max_dp": float(np.max(sched[reduced])) if reduced.any() else 0.0,
        "trigger_count": len(triggers),
        "recover_count": len(recovers),
        "cut_energy_pos": cut_energy_pos,
        "cut_energy_neg": cut_energy_neg,
        "rev_loss_pos": rev_loss_pos,
        "rev_save_neg": rev_save_neg,
        "net_rev": net_rev,
    }


# ---------------- 汇总 ----------------
def run_analysis(source=DATA_PATH, capacities=None):
    """按 capacities({sheet: (显示名, 额定容量MW)}) 逐场站×7场景回放，返回
    (rows, event_rows, recov_rows, adj_rows, full_rows) 五元组(内存态，不落盘)。"""
    if capacities is None:
        capacities = STATIONS
    all_data = load_sheets(source, list(capacities.keys()))
    rows = []
    event_rows = []
    recov_rows = []    # 恢复事件(用于Excel恢复事件明细)
    adj_rows = []      # 仅"被调整"周期(用于Excel调整明细)
    full_rows = []     # 全量时间序列(用于CSV全量明细)

    for sheet, (real_name, rated) in capacities.items():
        d = all_data[sheet]
        for sc_id, sc in SCENARIOS.items():
            r = simulate(d, rated, sc)
            sched = r["sched"]
            rows.append({
                "场站": sheet, "实际场站": real_name, "装机容量MW": rated,
                "场景": sc_id, "触发比例": sc["trigger_ratio"], "α": sc["alpha"],
                "β": sc["beta"], "γ": sc["gamma"],
                "触发次数": r["trigger_count"], "恢复次数": r["recover_count"],
                "压减周期数": r["reduced_periods"],
                "压减电量MWh": round(r["cut_energy"], 2),
                "压减电量占比%": round(r["cut_energy_pct"], 3),
                "平均ΔP%": round(r["avg_dp"] * 100, 2),
                "最大ΔP%": round(r["max_dp"] * 100, 2),
                "压减出力峰值MW": round(r["cut_power_peak"], 2),
                "正价压减电量MWh": round(r["cut_energy_pos"], 2),
                "负价压减电量MWh": round(r["cut_energy_neg"], 2),
                "正价收益损失元": round(r["rev_loss_pos"], 2),
                "负价避免支出元": round(r["rev_save_neg"], 2),
                "净收益影响元": round(r["net_rev"], 2),
            })
            # 触发事件明细
            for (tidx, dP, dev0, dev1) in r["triggers"]:
                event_rows.append({
                    "场站": sheet, "实际场站": real_name, "场景": sc_id,
                    "触发时刻": d["ts"][tidx], "ΔP%": round(dP * 100, 2),
                    "偏差1%": round(dev0 * 100, 2), "偏差2%": round(dev1 * 100, 2),
                })
            # 恢复事件明细
            for (ridx, dP) in r["recovers"]:
                recov_rows.append({
                    "场站": sheet, "实际场站": real_name, "场景": sc_id,
                    "恢复时刻": d["ts"][ridx], "恢复前ΔP%": round(dP * 100, 2),
                })
            # 逐周期明细: 调整后计划值 + 压减后实际出力 + 偏差量比例 + 是否调整标记
            opt_plan = r["opt_plan"]
            new_actual = r["new_actual"]
            actual_arr = d["actual"]
            # 偏差量比例 = (计划-实际)/计划, 与触发判定口径一致(计划取调整后计划)
            with np.errstate(divide="ignore", invalid="ignore"):
                dev = np.where(opt_plan > 1e-6, (opt_plan - actual_arr) / opt_plan, 0.0)
            # 连续两个周期偏差量平均比例 (触发即取最近两个周期的该平均值×β)
            avg_dev = np.zeros(len(dev))
            avg_dev[1:] = (dev[:-1] + dev[1:]) / 2.0
            # 触发/恢复确认时刻索引(用于在时间序列上打标记)
            trig_dp = {tidx: dP for (tidx, dP, _, _) in r["triggers"]}
            recov_idx = {ridx for (ridx, _) in r["recovers"]}
            for t in range(len(d["forecast"])):
                f = float(d["forecast"][t])
                p = float(d["plan"][t])
                a = float(actual_arr[t])
                adj = sched[t] > 0
                cut = float(np.maximum(0.0, a - opt_plan[t])) if adj else 0.0
                price_t = float(d["price"][t])
                price_ok = not np.isnan(price_t)
                rec = {
                    "场站": sheet, "实际场站": real_name, "场景": sc_id,
                    "时间": d["ts"][t],
                    "原上报预测MW": round(f, 3),
                    "原出清计划MW": round(p, 3),
                    "调整后计划MW": round(float(opt_plan[t]), 3),
                    "是否调整": "是" if adj else "否",
                    "触发标记": f"触发ΔP{trig_dp[t]*100:.2f}%" if t in trig_dp else "",
                    "恢复标记": "恢复" if t in recov_idx else "",
                    "ΔP%": round(float(sched[t]) * 100, 2) if adj else 0.0,
                    "之前实际出力MW": round(a, 3),
                    "压减后实际出力MW": round(float(new_actual[t]), 3),
                    "压减量MW": round(cut, 3),
                    "实时电价元/MWh": round(price_t, 2) if price_ok else None,
                    "收益影响元": round(cut * DT_HOUR * price_t, 2) if price_ok else None,
                    "偏差量比例%": round(float(dev[t]) * 100, 2),
                    "连续两周期平均偏差比例%": round(float(avg_dev[t]) * 100, 2),
                }
                full_rows.append(rec)
                if adj:
                    adj_rows.append(rec)
    return rows, event_rows, recov_rows, adj_rows, full_rows


def run_all():
    """CLI 入口：用默认数据文件与内置 STATIONS 跑全量分析。"""
    return run_analysis(DATA_PATH, STATIONS)


def build_workbook(rows, event_rows, recov_rows, adj_rows, full_rows):
    """构建汇总 Excel 工作簿(不落盘)，供 CLI 与网页复用。"""
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")

    def style_header(ws):
        for c in ws[1]:
            c.font = bold
            c.fill = fill

    # Sheet1 汇总对比
    ws1 = wb.active
    ws1.title = "汇总对比"
    header = ["场站", "实际场站", "装机容量MW", "场景", "触发比例", "α", "β", "γ",
              "触发次数", "恢复次数", "压减周期数", "压减电量MWh", "压减电量占比%",
              "平均ΔP%", "最大ΔP%", "压减出力峰值MW",
              "正价压减电量MWh", "负价压减电量MWh", "正价收益损失元", "负价避免支出元", "净收益影响元"]
    ws1.append(header)
    for r in rows:
        ws1.append([r[h] for h in header])
    style_header(ws1)

    # Sheet2 触发事件明细
    ws2 = wb.create_sheet("触发事件明细")
    eh = ["场站", "实际场站", "场景", "触发时刻", "ΔP%", "偏差1%", "偏差2%"]
    ws2.append(eh)
    for r in event_rows:
        ws2.append([r[h] for h in eh])
    style_header(ws2)

    # Sheet2b 恢复事件明细
    ws5 = wb.create_sheet("恢复事件明细")
    rh = ["场站", "实际场站", "场景", "恢复时刻", "恢复前ΔP%"]
    ws5.append(rh)
    for r in recov_rows:
        ws5.append([r[h] for h in rh])
    style_header(ws5)

    # Sheet3 调整明细(仅被调整周期)
    ws4 = wb.create_sheet("调整明细")
    dh = ["场站", "实际场站", "场景", "时间", "原上报预测MW", "原出清计划MW",
          "调整后计划MW", "是否调整", "触发标记", "恢复标记", "ΔP%", "之前实际出力MW", "压减后实际出力MW", "压减量MW",
          "实时电价元/MWh", "收益影响元", "偏差量比例%", "连续两周期平均偏差比例%"]
    ws4.append(dh)
    for r in adj_rows:
        ws4.append([r[h] for h in dh])
    style_header(ws4)

    # Sheet4 参数与口径
    ws3 = wb.create_sheet("参数与口径")
    notes = [
        ["项", "说明"],
        ["出清周期", "15 分钟"],
        ["触发门槛", "计划值 > 额定容量×10% 且连续2周期偏差 > 触发比例；窗口内计划值=调整后计划=预测×(1-ΔP%)(原始出清计划已被优化取代)"],
        ["计算延迟", "触发确认后跳过1个周期，自第+2周期起连续压减α个周期"],
        ["优化公式", "ΔP% = β × ((P计划,-1-P实际,-1)/P计划,-1 + (P计划,-2-P实际,-2)/P计划,-2)/2"],
        ["压减对象", "超短期预测功率(第2点)，P优化 = P原上报 × (1-ΔP%)"],
        ["恢复逻辑", "连续2周期实际≥计划后分γ周期恢复(第1次ΔP%/2，第2次全额)"],
        ["调整后计划值", "被调整周期=超短期预测×(1-ΔP%)；未调整周期=原出清计划"],
        ["压减后实际出力", "min(之前实际出力, 压减后计划)，未调整周期保持原实际出力"],
        ["压减量/压减电量", "max(0, 之前实际出力-压减后计划)×0.25h 累加，即实际出力被真正削减的电量"],
        ["收益口径", "收益影响=压减电量(MWh)×实时电价(元/MWh)；正价时段=少赚(损失)，负价时段=少亏(避免支出)；净收益影响=正价损失-负价避免，正=净少赚"],
        ["是否调整标记", "ΔP%>0 即标记为'是'"],
        ["触发/恢复标记", "全量明细中触发确认周期标'触发'、恢复确认周期标'恢复'；恢复事件另见'恢复事件明细'表"],
        ["场站映射", "玉山=临沭龙源(98.2MW)、东盘=玉山龙源(76.5MW)；柳格龙源(19.2MW)无数据未纳入"],
    ]
    for row in notes:
        ws3.append(row)
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 80
    style_header(ws3)

    return wb


def build_csv(full_rows):
    """构建全量明细 CSV 字符串(utf-8，调用方可自行加 BOM 或写文件)。"""
    import csv
    import io
    fh = ["场站", "实际场站", "场景", "时间", "原上报预测MW", "原出清计划MW",
          "调整后计划MW", "是否调整", "触发标记", "恢复标记", "ΔP%", "之前实际出力MW", "压减后实际出力MW", "压减量MW",
          "实时电价元/MWh", "收益影响元", "偏差量比例%", "连续两周期平均偏差比例%"]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fh)
    w.writeheader()
    for r in full_rows:
        w.writerow(r)
    return buf.getvalue()


def write_report(rows, event_rows, recov_rows, adj_rows, full_rows):
    """CLI 输出：写 Excel/CSV/报告到 OUT_DIR，返回文件路径。"""
    os.makedirs(OUT_DIR, exist_ok=True)

    wb = build_workbook(rows, event_rows, recov_rows, adj_rows, full_rows)
    xlsx_path = os.path.join(OUT_DIR, "汇总结果.xlsx")
    wb.save(xlsx_path)

    csv_path = os.path.join(OUT_DIR, "调整后计划_全量明细.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        f.write(build_csv(full_rows))

    # ---------- Markdown 报告 ----------
    md = ["# 新能源发电能力优化 场站影响测算报告\n",
          "> 依据《山东电网新能源发电能力优化试运行方案》7个场景，对8个场站的历史数据回放优化逻辑。\n"]
    md.append("## 一、各场景参数\n")
    md.append("| 场景 | 触发比例 | α(优化周期) | β(修正系数) | γ(恢复周期) |")
    md.append("|---|---|---|---|---|")
    for sid, sc in SCENARIOS.items():
        md.append(f"| {sid} | {sc['trigger_ratio']} | {sc['alpha']} | {sc['beta']} | {sc['gamma']} |")

    md.append("\n## 二、场站×场景 汇总（压减电量与触发动态）\n")
    md.append("| 场站 | 场景 | 触发次数 | 恢复次数 | 压减周期数 | 压减电量(MWh) | 占实际电量% | 平均ΔP% | 最大ΔP% |")
    md.append("|---|---|---|---|---|---|---|---|---|")
    for r in rows:
        md.append(f"| {r['场站']} | {r['场景']} | {r['触发次数']} | {r['恢复次数']} | "
                  f"{r['压减周期数']} | {r['压减电量MWh']} | {r['压减电量占比%']} | "
                  f"{r['平均ΔP%']} | {r['最大ΔP%']} |")

    # 按场景汇总
    md.append("\n## 三、按场景汇总（8场站合计）\n")
    md.append("| 场景 | 合计触发次数 | 合计压减电量(MWh) | 净收益影响(元) |")
    md.append("|---|---|---|---|")
    per_sc = {}
    for r in rows:
        per_sc.setdefault(r["场景"], [0, 0.0, 0.0])
        per_sc[r["场景"]][0] += r["触发次数"]
        per_sc[r["场景"]][1] += r["压减电量MWh"]
        per_sc[r["场景"]][2] += r["净收益影响元"]
    for sid in sorted(per_sc):
        md.append(f"| {sid} | {per_sc[sid][0]} | {round(per_sc[sid][1], 2)} | {round(per_sc[sid][2], 2)} |")

    # 收益分析
    md.append("\n## 四、收益影响（压减电量 × 实时电价）\n")
    md.append("> 正价时段少发=少赚(损失)；负价时段少发=少亏(避免支出)；净收益影响=正价损失-负价避免，正=净少赚，负=净少亏。\n")
    md.append("| 场站 | 场景 | 正价压减电量(MWh) | 负价压减电量(MWh) | 正价收益损失(元) | 负价避免支出(元) | 净收益影响(元) |")
    md.append("|---|---|---|---|---|---|---|")
    for r in rows:
        md.append(f"| {r['实际场站']} | {r['场景']} | {r['正价压减电量MWh']} | {r['负价压减电量MWh']} | "
                  f"{r['正价收益损失元']} | {r['负价避免支出元']} | {r['净收益影响元']} |")

    md_path = os.path.join(OUT_DIR, "报告.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    return xlsx_path, md_path, csv_path


# ---------------- 参考算例自检 ----------------
def selfcheck():
    """用附件2参考算例验证公式: 预测/计划100MW、实际80MW、β=1.5 -> 优化值70MW。"""
    n = 12
    forecast = np.full(n, 100.0)
    plan = np.full(n, 100.0)
    # T1/T2 实际80(偏差20%触发), 其余实际100(执行到位, 走恢复)
    actual = np.full(n, 100.0)
    actual[0] = actual[1] = 80.0
    d = {"ts": np.array([f"T{i}" for i in range(n)]), "forecast": forecast,
         "plan": plan, "actual": actual, "avail": np.full(n, 100.0),
         "price": np.full(n, 300.0)}
    sc = dict(trigger_ratio=0.1, alpha=4, beta=1.5, gamma=2)
    r = simulate(d, 100.0, sc)
    sched = r["sched"]
    # T1=idx0,T2=idx1 触发 -> 自 idx3(T4) 起压减, ΔP%=β*20%=30% -> 100*0.7=70MW
    opt_val = forecast[3] * (1 - sched[3])
    # 附件2算例: T4~T7 四个周期完整压减(30%)，压减窗口内不恢复；
    # 执行到位后 T8(idx7)=ΔP%/2=15%、T9(idx8)=0 分两批恢复
    window_ok = all(abs(sched[k] - 0.30) < 1e-9 for k in (3, 4, 5, 6))
    recov_ok = abs(sched[7] - 0.15) < 1e-9 and abs(sched[8]) < 1e-9
    ok = window_ok and recov_ok and r["trigger_count"] == 1
    print(f"[自检] T4优化值(期望70MW) = {opt_val:.1f} MW, 压减窗口30%x4={'通过' if window_ok else '失败'}, "
          f"恢复15%/0={'通过' if recov_ok else '失败'}, 触发次数={r['trigger_count']}, "
          f"校验={'通过' if ok else '未通过'}")
    return ok and abs(opt_val - 70) < 1e-6


if __name__ == "__main__":
    assert selfcheck(), "参考算例自检未通过！"
    rows, events, recov_rows, adj_rows, full_rows = run_all()
    xlsx, md, csv_path = write_report(rows, events, recov_rows, adj_rows, full_rows)
    print(f"完成。汇总: {xlsx}")
    print(f"完成。全量明细: {csv_path}")
    print(f"完成。报告: {md}")
